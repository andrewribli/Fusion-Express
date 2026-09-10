"""Fill missing catalog prices from ParknShop (pns.hk).

Reads Product List Catalog.xlsx rows where Price (HKD) is blank, searches
pns.hk with Playwright, and writes results to parknshop-enrichment.json.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from enrich_from_parknshop import (  # noqa: E402
    CACHE,
    PNS_TIMEOUT_MS,
    STEALTH_JS,
    UA,
    XLSX,
    best_hit,
    extract_hits,
    search_query,
)
import urllib.parse  # noqa: E402


def search_pns_prices(page, query: str) -> list[dict]:
    url = "https://www.pns.hk/en/search?q=" + urllib.parse.quote(query)
    page.goto(url, wait_until="domcontentloaded", timeout=45000)
    page.wait_for_timeout(2500)
    try:
        page.wait_for_selector('a[href*="/p/"]', timeout=12000)
    except Exception:
        return []
    hits = extract_hits(page, require_image=False)
    if hits:
        return hits
    short = " ".join(query.split()[:5])
    if short != query:
        url = "https://www.pns.hk/en/search?q=" + urllib.parse.quote(short)
        page.goto(url, wait_until="domcontentloaded", timeout=45000)
        page.wait_for_timeout(2500)
        hits = extract_hits(page, require_image=False)
    return hits
from sanitize_product_text import clean_brand, is_unknown  # noqa: E402

ENRICHMENT = CACHE
FOODPANDA_PRICE: dict[str, float] = {}


def load_foodpanda_prices() -> dict[str, float]:
    if FOODPANDA_PRICE:
        return FOODPANDA_PRICE
    for path in (
        ROOT / "FoodPanda Fusion Data.txt",
        ROOT / "FoodPanda Data2.txt",
    ):
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            parts = line.split("\t")
            if len(parts) < 3:
                continue
            name = parts[1].strip().strip("'").strip('"')
            price_raw = parts[2].strip().strip("'").strip('"')
            if name and re.match(r"^\d+(?:\.\d+)?$", price_raw):
                FOODPANDA_PRICE[name.lower()] = float(price_raw)
    return FOODPANDA_PRICE


def fuzzy_fp_price(name: str, brand: str | None, pool: dict[str, float]) -> float | None:
    key = name.lower()
    if key in pool:
        return pool[key]
    q = f"{brand or ''} {name}".lower()
    qtok = {t for t in re.findall(r"[a-z0-9]{3,}", q)}
    best = None
    best_score = 0.0
    for other, price in pool.items():
        otok = {t for t in re.findall(r"[a-z0-9]{3,}", other)}
        if not qtok or not otok:
            continue
        score = len(qtok & otok) / max(len(qtok), 1)
        if score > best_score:
            best_score = score
            best = price
    return best if best_score >= 0.55 else None


def price_from_product_page(page, href: str) -> float | None:
    try:
        page.goto(href, wait_until="domcontentloaded", timeout=PNS_TIMEOUT_MS)
        page.wait_for_timeout(1500)
        body = page.locator("body").inner_text(timeout=5000)
        matches = re.findall(r"\$([0-9]+(?:\.[0-9]{2})?)", body)
        for raw in matches:
            val = float(raw)
            if 1 <= val <= 5000:
                return val
    except Exception:
        return None
    return None


def load_missing_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    missing: list[dict] = []
    for row in rows[1:]:
        if not row or len(row) < 6 or not row[2]:
            continue
        price = row[3]
        if price is not None and float(price) > 0:
            continue
        raw_name = str(row[2]).strip()
        if re.search(
            r"unknown|unclear|partial|facings?|right-edge|left-edge|bottom-edge|"
            r"product unknown|variant unknown|obscured",
            raw_name,
            re.I,
        ):
            continue
        raw_brand = str(row[1]).strip() if row[1] else None
        missing.append(
            {
                "rawName": raw_name,
                "brand": clean_brand(raw_brand) if not is_unknown(raw_brand) else None,
                "rawBrand": raw_brand,
            }
        )
    return missing


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--delay", type=float, default=0.6)
    args = parser.parse_args()

    cache: dict[str, dict] = {}
    if ENRICHMENT.exists():
        cache = json.loads(ENRICHMENT.read_text(encoding="utf-8"))

    todo = []
    for row in load_missing_rows():
        key = row["rawName"]
        if cache.get(key, {}).get("price") is not None:
            continue
        todo.append(row)
    if args.limit:
        todo = todo[: args.limit]

    print(f"Filling prices for {len(todo)} products ({len(cache)} cached)")

    stats = {"pns": 0, "miss": 0}

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=False,
            args=["--disable-blink-features=AutomationControlled", "--no-sandbox"],
        )
        context = browser.new_context(
            user_agent=UA,
            locale="en-HK",
            viewport={"width": 1365, "height": 900},
        )
        context.add_init_script(STEALTH_JS)
        page = context.new_page()
        page.set_default_timeout(PNS_TIMEOUT_MS)

        for i, row in enumerate(todo, 1):
            raw_name = row["rawName"]
            brand = row["brand"]
            query = search_query(row.get("rawBrand"), raw_name)
            entry = cache.get(raw_name, {})
            price = None
            source = ""

            try:
                hits = search_pns_prices(page, query)
                hit = best_hit(query, brand, hits)
                if hit:
                    price = hit.get("price")
                    if price is None and hit.get("href"):
                        price = price_from_product_page(page, hit["href"])
                    if price is not None:
                        entry["pnsName"] = hit.get("name")
                        entry["pnsUrl"] = hit.get("href")
                        source = "pns"
            except Exception as exc:
                print(f"[{i}/{len(todo)}] ERR {raw_name}: {exc}")

            if price is not None:
                entry["price"] = price
                cache[raw_name] = entry
                stats[source or "pns"] = stats.get(source or "pns", 0) + 1
                print(f"[{i}/{len(todo)}] {source.upper()} ${price} {raw_name[:55]}")
            else:
                stats["miss"] += 1
                print(f"[{i}/{len(todo)}] MISS {raw_name} ({query})")

            if i % 10 == 0:
                ENRICHMENT.write_text(
                    json.dumps(cache, ensure_ascii=False, indent=2) + "\n",
                    encoding="utf-8",
                )
            time.sleep(args.delay)

        browser.close()

    ENRICHMENT.write_text(json.dumps(cache, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # Patch Excel if not locked.
    if XLSX.exists():
        try:
            wb = openpyxl.load_workbook(XLSX)
            ws = wb["Sheet1"]
            for r in ws.iter_rows(min_row=2):
                name = r[2].value
                if not name:
                    continue
                hit = cache.get(str(name).strip())
                if hit and hit.get("price") is not None and r[3].value is None:
                    r[3].value = hit["price"]
            wb.save(XLSX)
        except PermissionError:
            print(f"Warning: could not save {XLSX} (file may be open)")

    print("done", stats)


if __name__ == "__main__":
    main()

"""Build foodpanda-fusion-catalog.json from Product List Catalog.xlsx."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

from sanitize_product_text import clean_brand, clean_name

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Product List Catalog.xlsx"
OUT = ROOT / "packages" / "shared" / "data" / "foodpanda-fusion-catalog.json"
OLD_CATALOG = OUT
ENRICHMENT = ROOT / "packages" / "shared" / "data" / "parknshop-enrichment.json"
DRIVE_IMAGES = ROOT / "apps" / "web" / "public" / "images" / "catalog"

# Photo-notes / placeholders we could not identify on ParknShop.
UNIDENTIFIED_RE = re.compile(
    r"unknown|unclear|partial|facings?|right-edge|left-edge|bottom-edge|"
    r"product unknown|variant unknown|obscured",
    re.I,
)


def should_skip_row(raw_name: str, excel_price: object, enriched: dict) -> bool:
    """Drop items we guessed instead of identifying on ParknShop."""
    if UNIDENTIFIED_RE.search(raw_name):
        return True
    guessed = enriched.get("priceSource") in {"median", "foodpanda"}
    if guessed:
        return True
    no_excel = excel_price is None or (isinstance(excel_price, (int, float)) and float(excel_price) <= 0)
    if no_excel and not enriched.get("pnsUrl") and not enriched.get("pnsName"):
        if enriched.get("price") is not None:
            return True
    return False

SECTION_ORDER = ("Groceries", "Fresh Food")

# Meat aisles: beef first, then typical campus demand.
MEAT_SUB_ORDER = ("Beef", "Chicken", "Pork", "Seafood", "Others")

GROCERIES_ALIASES = {
    "groceries",
    "non-refrigerated",
    "non refrigerated",
    "nonrefrigerated",
    "dry",
    "dry goods",
}
FRESH_FOOD_ALIASES = {
    "fresh food",
    "refrigerated",
    "chilled",
    "chilled & refrigerated",
}


def normalize_section(raw: object) -> str:
    label = str(raw).strip() if raw else "Groceries"
    key = " ".join(label.lower().replace("_", " ").replace("-", " ").split())
    compact = key.replace(" ", "")
    if key in GROCERIES_ALIASES or compact in {"nonrefrigerated", "drygoods"}:
        return "Groceries"
    if key in FRESH_FOOD_ALIASES or compact in {"freshfood", "refrigerated"}:
        return "Fresh Food"
    return label or "Groceries"


CHILLED_GROCERY_SUBS = {
    "ice cream": "Ice Cream",
    "frozen vegetables": "Frozen Vegetables",
    "frozen meat": "Frozen Meat",
    "meat": "Meat",
}


def normalize_subcategory(raw: object) -> str:
    label = str(raw).strip() if raw else ""
    return label or ""


def infer_fresh_subcategory(name: str, brand: str | None) -> str:
    n = f"{brand or ''} {name}".lower()
    if any(w in n for w in ("sandwich", "onigiri", "bento", "ready meal")):
        return "Ready Meals"
    if any(w in n for w in ("drink", "bottled")):
        return "Chilled Drinks"
    if "salmon" in n or "seafood" in n:
        return "Seafood"
    if any(w in n for w in ("tofu", "bean curd")):
        return "Dairy"
    if "kimchi" in n:
        return "Vegetables"
    if "egg" in n:
        return "Dairy"
    if any(
        w in n
        for w in (
            "yogurt",
            "yoghurt",
            "cream",
            "butter",
            "cheese",
            "cheddar",
            "mozzarella",
            "parmesan",
            "spread",
            "dairy",
        )
    ):
        return "Dairy"
    if any(
        w in n
        for w in (
            "blueberry",
            "blueberries",
            "kiwi",
            "fruit",
            "watermelon",
            "melon",
            "grape",
            "peach",
            "citrus",
            "orange",
        )
    ):
        return "Fruit"
    return "Other"


def classify_meat(name: str, brand: str | None) -> str:
    n = f"{brand or ''} {name}".lower()
    if any(
        w in n
        for w in ("plant-based", "plant based", "meat zero", "vegetarian", "vegan")
    ):
        return "Others"
    if any(
        w in n
        for w in (
            "salmon",
            "prawn",
            "shrimp",
            "fish",
            "seafood",
            "abalone",
            "duck leg",
            "confit duck",
        )
    ):
        return "Seafood"
    if any(
        w in n
        for w in ("chicken", "poultry", "wing", "thigh", "breast", "drumstick")
    ):
        return "Chicken"
    if any(
        w in n
        for w in (
            "beef",
            "steak",
            "striploin",
            "angus",
            "grain fed",
            "hot pot beef",
            "halal beef",
        )
    ):
        return "Beef"
    if any(
        w in n
        for w in (
            "pork",
            "belly",
            "spare rib",
            "sparerib",
            "mince",
            "ground pork",
            "collar",
            "sausage",
            "bacon",
            "ham",
            "patty",
        )
    ):
        return "Pork"
    return "Others"


def resolve_section_and_sub(
    raw_section: object, raw_sub: object, name: str, brand: str | None
) -> tuple[str, str]:
    section = normalize_section(raw_section)
    sub = normalize_subcategory(raw_sub)

    # Frozen / chilled grocery rows belong in Fresh Food.
    if section == "Groceries" and sub.lower() in CHILLED_GROCERY_SUBS:
        section = "Fresh Food"
        sub = CHILLED_GROCERY_SUBS[sub.lower()]

    meat_subs = {"meat", "frozen meat", "pork", "beef", "chicken", "seafood", "others"}
    if sub.lower() in meat_subs:
        classified = classify_meat(name, brand)
        n = f"{brand or ''} {name}".lower()
        if sub.lower() in {"meat", "frozen meat"}:
            sub = classified
        elif classified == "Others" and any(
            w in n for w in ("plant-based", "plant based", "meat zero")
        ):
            sub = "Others"
        elif classified in MEAT_SUB_ORDER and classified != sub:
            if any(w in n for w in ("beef", "chicken", "pork", "seafood")):
                sub = classified

    if section == "Fresh Food" and not sub:
        sub = infer_fresh_subcategory(name, brand)

    if not sub:
        sub = "Other"
    return section, sub


def parse_weight(raw: object) -> float | None:
    if raw is None:
        return None
    m = re.search(r"(\d+(?:\.\d+)?)\s*(kg|g|ml|l)\b", str(raw).strip(), re.I)
    if not m:
        return None
    qty = float(m.group(1))
    unit = m.group(2).lower()
    if unit == "g":
        return round(qty / 1000, 3)
    if unit == "kg":
        return round(qty, 3)
    if unit == "ml":
        return round(qty / 1000, 3)
    if unit == "l":
        return round(qty, 3)
    return None


def parse_multibuy(promo: object) -> tuple[float, int] | None:
    if promo is None:
        return None
    m = re.match(r"(\d+(?:\.\d+)?)\s*/\s*(\d+)", str(promo).strip())
    if not m:
        return None
    return float(m.group(1)), int(m.group(2))


def load_old_images() -> dict[str, str]:
    if not OLD_CATALOG.exists():
        return {}
    data = json.loads(OLD_CATALOG.read_text(encoding="utf-8"))
    images: dict[str, str] = {}

    def walk(nodes: list) -> None:
        for node in nodes:
            for item in node.get("items") or []:
                img = item.get("image")
                name = (item.get("name") or "").strip().lower()
                if img and name:
                    images[name] = img
            walk(node.get("subcategories") or [])

    walk(data.get("categories") or [])
    return images


def match_image(name: str, images: dict[str, str]) -> str | None:
    key = name.strip().lower()
    if key in images:
        return images[key]
    for old_name, url in images.items():
        if key in old_name or old_name in key:
            return url
    return None


def drive_image_url(raw: object) -> str | None:
    """Map Excel Image Source (e.g. IMG_7053.HEIC) to a local catalog photo."""
    if not raw:
        return None
    name = str(raw).strip()
    if not name.upper().endswith(".HEIC"):
        return None
    stem = Path(name).stem.lower().replace("_", "-")
    if (DRIVE_IMAGES / f"{stem}.jpg").exists():
        return f"/images/catalog/{stem}.jpg"
    return None


def load_enrichment() -> dict[str, dict]:
    if not ENRICHMENT.exists():
        return {}
    return json.loads(ENRICHMENT.read_text(encoding="utf-8"))


def main() -> None:
    # Capture images from previous catalog before overwrite.
    images = load_old_images()
    enrichment = load_enrichment()

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))

    buckets: dict[str, dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    matched = 0

    for row in rows[1:]:
        if not row or len(row) < 6 or not row[2]:
            continue
        raw_name = str(row[2]).strip()
        raw_brand = str(row[1]).strip() if row[1] else None
        price = row[3]
        enriched = enrichment.get(raw_name) or {}
        if should_skip_row(raw_name, price, enriched):
            continue
        brand = clean_brand(raw_brand)
        category, subcategory = resolve_section_and_sub(
            row[4] if len(row) > 4 else None,
            row[5] if len(row) > 5 else None,
            raw_name,
            brand or raw_brand,
        )
        use_pns_name = bool(UNIDENTIFIED_RE.search(raw_name))
        name = enriched.get("resolvedName") or clean_name(
            raw_name,
            brand,
            subcategory,
            pns_name=enriched.get("pnsName") if use_pns_name else None,
        )
        weight = parse_weight(row[6] if len(row) > 6 else None)
        multibuy = parse_multibuy(row[7] if len(row) > 7 else None)
        if enriched.get("subcategory"):
            subcategory = enriched["subcategory"]
        drive_img = drive_image_url(row[11] if len(row) > 11 else None)
        image = enriched.get("image") or drive_img or match_image(name, images)
        if image:
            matched += 1
        resolved_price = price
        if resolved_price is None and enriched.get("price") is not None:
            resolved_price = enriched["price"]

        item: dict = {
            "name": name,
            "price": float(resolved_price) if resolved_price is not None else 0.0,
            "category": category,
            "subcategory": subcategory,
        }
        if brand:
            item["brand"] = brand
        elif enriched.get("pnsBrand"):
            item["brand"] = enriched["pnsBrand"]
        if image:
            item["image"] = image
        if weight is not None:
            item["weight"] = weight
        if multibuy is not None:
            item["bulkDealPrice"] = multibuy[0]
            item["bulkDealQty"] = multibuy[1]

        buckets[category][subcategory].append(item)

    catalog = {"categories": []}
    for cat_name in SECTION_ORDER:
        subs = buckets.get(cat_name, {})
        node = {"name": cat_name, "subcategories": []}

        def sub_key(s: str, counts: dict[str, list] = subs) -> tuple:
            if cat_name == "Fresh Food" and s in MEAT_SUB_ORDER:
                return (0, MEAT_SUB_ORDER.index(s))
            return (1, -len(counts[s]), s)

        for sub_name in sorted(subs.keys(), key=sub_key):
            node["subcategories"].append(
                {"name": sub_name, "items": subs[sub_name]}
            )
        catalog["categories"].append(node)

    # Keep any unexpected top-level categories from the sheet.
    for cat_name, subs in sorted(buckets.items()):
        if cat_name in SECTION_ORDER:
            continue
        node = {"name": cat_name, "subcategories": []}
        for sub_name in sorted(subs.keys(), key=lambda s: (-len(subs[s]), s)):
            node["subcategories"].append(
                {"name": sub_name, "items": subs[sub_name]}
            )
        catalog["categories"].append(node)

    OUT.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    total = sum(len(v) for c in buckets.values() for v in c.values())
    print(f"Wrote {OUT} ({total} items, {matched} with images)")


if __name__ == "__main__":
    main()
